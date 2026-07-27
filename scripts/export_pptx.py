#!/usr/bin/env python3
"""
PPTX Export Script for Nature Skills Workflow.

Generates PowerPoint presentations from JSON or Markdown outlines,
with support for speaker notes, image embedding, and chart placeholders.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)


@dataclass
class SlideContent:
    """Container for a single slide's content."""

    title: str
    content: List[str] = field(default_factory=list)
    notes: Optional[str] = None
    image_path: Optional[Path] = None
    image_caption: Optional[str] = None
    layout: str = "title_and_content"  # title, title_and_content, section_header, two_content, blank
    chart_placeholder: bool = False


class OutlineParser:
    """Parse JSON or Markdown outlines into slide structures."""

    @staticmethod
    def parse_json(filepath: Path) -> List[SlideContent]:
        """Parse JSON outline file.

        Expected JSON format:
        {
            "slides": [
                {
                    "title": "Slide Title",
                    "content": ["Bullet 1", "Bullet 2"],
                    "notes": "Speaker notes",
                    "image": "path/to/image.png",
                    "layout": "title_and_content"
                }
            ]
        }
        """
        try:
            data = json.loads(filepath.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            logger.error(f"Invalid JSON in {filepath}: {e}")
            return []

        slides_data = data.get("slides", [])
        if not slides_data:
            logger.warning(f"No slides found in {filepath}")
            return []

        slides = []
        for i, slide_data in enumerate(slides_data):
            slide = SlideContent(
                title=slide_data.get("title", f"Slide {i + 1}"),
                content=slide_data.get("content", []),
                notes=slide_data.get("notes"),
                image_path=Path(slide_data["image"]) if slide_data.get("image") else None,
                image_caption=slide_data.get("image_caption"),
                layout=slide_data.get("layout", "title_and_content"),
                chart_placeholder=slide_data.get("chart_placeholder", False),
            )
            slides.append(slide)

        logger.info(f"Parsed {len(slides)} slides from JSON: {filepath}")
        return slides

    @staticmethod
    def parse_markdown(filepath: Path) -> List[SlideContent]:
        """Parse Markdown outline file.

        Expected Markdown format:
        # Slide Title
        - Bullet point 1
        - Bullet point 2

        Notes: Speaker notes here

        Image: path/to/image.png
        Caption: Image caption

        ---

        # Next Slide
        ...
        """
        try:
            content = filepath.read_text(encoding="utf-8")
        except Exception as e:
            logger.error(f"Failed to read {filepath}: {e}")
            return []

        slides = []
        current_slide: Optional[Dict[str, Any]] = None

        lines = content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            # Slide separator
            if line == "---":
                if current_slide:
                    slides.append(SlideContent(**current_slide))
                    current_slide = None
                i += 1
                continue

            # New slide title
            if line.startswith("# ") and not line.startswith("## "):
                if current_slide:
                    slides.append(SlideContent(**current_slide))
                current_slide = {
                    "title": line[2:].strip(),
                    "content": [],
                    "notes": None,
                    "image_path": None,
                    "image_caption": None,
                    "layout": "title_and_content",
                    "chart_placeholder": False,
                }
                i += 1
                continue

            if current_slide is None:
                i += 1
                continue

            # Bullet points
            if line.startswith("- ") or line.startswith("* "):
                current_slide["content"].append(line[2:].strip())
                i += 1
                continue

            # Notes
            if line.lower().startswith("notes:"):
                current_slide["notes"] = line[6:].strip()
                i += 1
                continue

            # Image
            if line.lower().startswith("image:"):
                img_path = line[6:].strip()
                current_slide["image_path"] = Path(img_path)
                i += 1
                continue

            # Image caption
            if line.lower().startswith("caption:"):
                current_slide["image_caption"] = line[8:].strip()
                i += 1
                continue

            # Layout
            if line.lower().startswith("layout:"):
                current_slide["layout"] = line[7:].strip().lower()
                i += 1
                continue

            # Chart placeholder
            if line.lower().startswith("chart:"):
                current_slide["chart_placeholder"] = True
                i += 1
                continue

            # Multi-line content continuation
            if line and not line.startswith("#"):
                current_slide["content"].append(line)

            i += 1

        # Add last slide
        if current_slide:
            slides.append(SlideContent(**current_slide))

        logger.info(f"Parsed {len(slides)} slides from Markdown: {filepath}")
        return slides


class PPTXGenerator:
    """Generate PowerPoint presentations."""

    # Layout mapping to python-pptx slide layouts
    LAYOUT_MAP = {
        "title": 0,           # Title Slide
        "title_and_content": 1,  # Title and Content
        "section_header": 2,     # Section Header
        "two_content": 3,        # Two Content
        "comparison": 4,         # Comparison
        "title_only": 5,         # Title Only
        "blank": 6,              # Blank
        "content_with_caption": 7,
        "picture_with_caption": 8,
    }

    def __init__(self, template_path: Optional[Path] = None):
        """Initialize generator.

        Args:
            template_path: Optional .pptx template file.
        """
        if template_path and template_path.exists():
            self.prs = Presentation(str(template_path))
            logger.info(f"Using template: {template_path}")
        else:
            self.prs = Presentation()
            # Set default 16:9 aspect ratio
            self.prs.slide_width = Inches(13.333)
            self.prs.slide_height = Inches(7.5)

        self._setup_styles()

    def _setup_styles(self) -> None:
        """Setup default text styles."""
        self.title_font_size = Pt(36)
        self.content_font_size = Pt(20)
        self.notes_font_size = Pt(14)
        self.title_color = RGBColor(0x1A, 0x1A, 0x2E)
        self.content_color = RGBColor(0x33, 0x33, 0x33)

    def _get_layout_index(self, layout_name: str) -> int:
        """Get slide layout index by name."""
        return self.LAYOUT_MAP.get(layout_name, 1)

    def add_slide(self, slide: SlideContent) -> None:
        """Add a single slide to the presentation.

        Args:
            slide: SlideContent object with slide data.
        """
        layout_idx = self._get_layout_index(slide.layout)
        try:
            slide_layout = self.prs.slide_layouts[layout_idx]
        except IndexError:
            logger.warning(f"Layout {layout_idx} not available, using default")
            slide_layout = self.prs.slide_layouts[1]

        pptx_slide = self.prs.slides.add_slide(slide_layout)

        # Set title
        if pptx_slide.shapes.title:
            title_frame = pptx_slide.shapes.title.text_frame
            title_frame.text = slide.title
            for paragraph in title_frame.paragraphs:
                paragraph.font.size = self.title_font_size
                paragraph.font.color.rgb = self.title_color
                paragraph.font.bold = True

        # Add content
        content_added = False
        for shape in pptx_slide.placeholders:
            if shape.placeholder_format.idx == 1:  # Content placeholder
                text_frame = shape.text_frame
                text_frame.word_wrap = True

                for i, bullet in enumerate(slide.content):
                    if i == 0:
                        p = text_frame.paragraphs[0]
                    else:
                        p = text_frame.add_paragraph()

                    p.text = bullet
                    p.font.size = self.content_font_size
                    p.font.color.rgb = self.content_color
                    p.level = 0
                    p.space_after = Pt(12)

                content_added = True
                break

        # If no content placeholder found, add textbox
        if not content_added and slide.content:
            left = Inches(0.5)
            top = Inches(1.5)
            width = Inches(12.333)
            height = Inches(5.5)
            txBox = pptx_slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.word_wrap = True

            for i, bullet in enumerate(slide.content):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text = bullet
                p.font.size = self.content_font_size
                p.font.color.rgb = self.content_color
                p.space_after = Pt(12)

        # Add image if specified
        if slide.image_path and slide.image_path.exists():
            try:
                # Position image on right side or center
                if slide.layout == "two_content":
                    left = Inches(7)
                    top = Inches(1.5)
                    width = Inches(5.5)
                else:
                    left = Inches(3)
                    top = Inches(2)
                    width = Inches(7)

                pptx_slide.shapes.add_picture(
                    str(slide.image_path),
                    left, top, width=width,
                )
                logger.info(f"  Added image: {slide.image_path}")

                # Add caption if provided
                if slide.image_caption:
                    caption_left = left
                    caption_top = top + Inches(4.5)
                    caption_box = pptx_slide.shapes.add_textbox(
                        caption_left, caption_top, width, Inches(0.5)
                    )
                    caption_tf = caption_box.text_frame
                    caption_tf.text = slide.image_caption
                    caption_tf.paragraphs[0].font.size = Pt(12)
                    caption_tf.paragraphs[0].font.italic = True
                    caption_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

            except Exception as e:
                logger.warning(f"  Failed to add image {slide.image_path}: {e}")

        # Add chart placeholder
        if slide.chart_placeholder:
            left = Inches(3)
            top = Inches(2)
            width = Inches(7)
            height = Inches(4)
            chart_box = pptx_slide.shapes.add_textbox(left, top, width, height)
            chart_tf = chart_box.text_frame
            chart_tf.text = "[CHART PLACEHOLDER]"
            chart_tf.paragraphs[0].font.size = Pt(24)
            chart_tf.paragraphs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            chart_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Add speaker notes
        if slide.notes:
            notes_slide = pptx_slide.notes_slide
            notes_tf = notes_slide.notes_text_frame
            notes_tf.text = slide.notes
            for paragraph in notes_tf.paragraphs:
                paragraph.font.size = self.notes_font_size

    def generate(self, slides: List[SlideContent], output_path: Path) -> bool:
        """Generate complete PPTX from slide list.

        Args:
            slides: List of SlideContent objects.
            output_path: Output .pptx file path.

        Returns:
            True if successful.
        """
        if not slides:
            logger.error("No slides to generate")
            return False

        logger.info(f"Generating PPTX with {len(slides)} slides...")

        for i, slide in enumerate(slides, 1):
            logger.info(f"  Slide {i}: {slide.title}")
            try:
                self.add_slide(slide)
            except Exception as e:
                logger.error(f"  Failed to add slide {i}: {e}")

        try:
            self.prs.save(str(output_path))
            size_kb = output_path.stat().st_size / 1024
            logger.info(f"PPTX saved: {output_path} ({size_kb:.1f} KB)")
            return True
        except Exception as e:
            logger.error(f"Failed to save PPTX: {e}")
            return False


def main() -> int:
    """Main entry point."""
    parser = argparse.ArgumentParser(
        description="Generate PowerPoint from JSON or Markdown outline",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  %(prog)s outline.json -o presentation.pptx
  %(prog)s outline.md -o slides.pptx
  %(prog)s outline.json --template template.pptx

Markdown outline format:
  # Slide Title
  - Bullet 1
  - Bullet 2
  Notes: Speaker notes
  Image: path/to/image.png
  ---
  # Next Slide
  ...
        """,
    )
    parser.add_argument(
        "input",
        type=Path,
        help="Input outline file (.json or .md)",
    )
    parser.add_argument(
        "-o", "--output",
        type=Path,
        required=True,
        help="Output PPTX file path",
    )
    parser.add_argument(
        "--template",
        type=Path,
        help="PPTX template file",
    )
    parser.add_argument(
        "-v", "--verbose",
        action="store_true",
        help="Enable verbose logging",
    )

    args = parser.parse_args()

    if args.verbose:
        logging.getLogger().setLevel(logging.DEBUG)

    if not args.input.exists():
        logger.error(f"Input file not found: {args.input}")
        return 1

    # Parse outline
    suffix = args.input.suffix.lower()
    if suffix == ".json":
        slides = OutlineParser.parse_json(args.input)
    elif suffix in (".md", ".markdown"):
        slides = OutlineParser.parse_markdown(args.input)
    else:
        logger.error(f"Unsupported input format: {suffix}. Use .json or .md")
        return 1

    if not slides:
        logger.error("No slides parsed from input")
        return 1

    # Generate PPTX
    generator = PPTXGenerator(template_path=args.template)
    success = generator.generate(slides, args.output)

    return 0 if success else 1


if __name__ == "__main__":
    sys.exit(main())
